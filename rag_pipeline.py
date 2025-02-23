from haystack.document_stores.in_memory import InMemoryDocumentStore

document_store = InMemoryDocumentStore()

from datasets import load_dataset
from haystack import Document
'''
bilgeyucel/seven-wonders

dataset = load_dataset("Solip-n/PiEGPT", split="train")
docs = [Document(content=doc["content"], meta=doc["meta"]) for doc in dataset]
'''
from haystack.document_stores.in_memory import InMemoryDocumentStore
from haystack import Document
import gdown
import os

url = "https://drive.google.com/drive/folders/17tD6iDYUa5ndfhSSE06cv16UaNszXNfD"
output_dir = "recipe_files"
gdown.download_folder(url, quiet=True, output=output_dir)

# Create documents from downloaded files
docs = []

for root, _, files in os.walk(output_dir):
    for file in files:
        file_path = os.path.join(root, file)
        
        # Handle different file types
        if file.endswith(".txt"):
            with open(file_path, "r") as f:
                content = f.read()
                
        elif file.endswith(".docx"):
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            content = "\n".join([para.text for para in doc.paragraphs])
            
        elif file.endswith(".pdf"):
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            content = "\n".join([page.extract_text() for page in reader.pages])
            
        else:
            continue  # Skip unsupported formats
            
        docs.append(
            Document(
                content=content,
                meta={"source": file, "path": file_path}
            )
        )

#pipeline setup
document_store = InMemoryDocumentStore()


from haystack.components.embedders import SentenceTransformersDocumentEmbedder

doc_embedder = SentenceTransformersDocumentEmbedder(model="sentence-transformers/all-MiniLM-L6-v2")
doc_embedder.warm_up()

docs_with_embeddings = doc_embedder.run(docs)
document_store.write_documents(docs_with_embeddings["documents"])

from haystack.components.embedders import SentenceTransformersTextEmbedder

text_embedder = SentenceTransformersTextEmbedder(model="sentence-transformers/all-MiniLM-L6-v2")

from haystack.components.retrievers.in_memory import InMemoryEmbeddingRetriever

retriever = InMemoryEmbeddingRetriever(document_store)

from haystack.components.builders import ChatPromptBuilder
from haystack.dataclasses import ChatMessage

template = [
    ChatMessage.from_user(
        """
Given the following information, answer the question.

Context:
{% for document in documents %}
    {{ document.content }}
{% endfor %}

Question: {{question}}
Answer:
"""
    )
]

prompt_builder = ChatPromptBuilder(template=template)

import os
from getpass import getpass
from haystack.utils import Secret
from haystack.components.generators.chat import OpenAIChatGenerator

if "TOGETHER_AI_API_KEY" not in os.environ:
    os.environ["TOGETHER_AI_API_KEY"] = "4a5416ae3c4d3befc622de5884ac0dbef1371aef12ae063d57eb54c9e6e10906"

chat_generator = OpenAIChatGenerator(model="deepseek-ai/DeepSeek-R1",
    api_key=Secret.from_env_var("TOGETHER_AI_API_KEY"),
    api_base_url="https://api.together.xyz/v1",
    streaming_callback=lambda chunk: print(chunk.content, end="", flush=True))


#PIPELINE
from haystack import Pipeline

basic_rag_pipeline = Pipeline()
# Add components to your pipeline
basic_rag_pipeline.add_component("text_embedder", text_embedder)
basic_rag_pipeline.add_component("retriever", retriever)
basic_rag_pipeline.add_component("prompt_builder", prompt_builder)
basic_rag_pipeline.add_component("llm", chat_generator)

# Now, connect the components to each other
basic_rag_pipeline.connect("text_embedder.embedding", "retriever.query_embedding")
basic_rag_pipeline.connect("retriever", "prompt_builder")
basic_rag_pipeline.connect("prompt_builder.prompt", "llm.messages")


def ask_question(question: str):
    response = basic_rag_pipeline.run({
        "text_embedder": {"text": question},
        "prompt_builder": {"question": question}
    })
    return response["llm"]["replies"][0].content


#question = "what is robotics"
#response = basic_rag_pipeline.run({"text_embedder": {"text": question}, "prompt_builder": {"question": question}})
#print(response["llm"]["replies"][0].text) #.content